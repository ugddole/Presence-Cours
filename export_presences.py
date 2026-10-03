"""Export Excel des présences UGD : un onglet par cours, une ligne par adhérent,
une colonne par séance (Présent / Absent / Excusé), totaux en formules.

Format attendu (construit par la route /admin/export.xlsx de app.py) :
    cours = [{
      "label":   "Baby gym – mercredi",
      "couleur": "#e53935",                                   # optionnel
      "seances": [{"cle": 12, "date": date(2026, 9, 9), "horaire": ""}, ...],
      "lignes":  [{"nom": "DUPONT", "prenom": "Léa", "essai": False,
                   "statuts": {12: "present"}}, ...],          # present | excuse | absent
    }, ...]
Une case vide = la personne ne figurait pas sur la liste ce jour-là (pas encore inscrite,
ou venue seulement en essai) ; elle n'est pas comptée dans son taux.
"""
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LIBELLES = {"present": "Présent", "excuse": "Excusé", "absent": "Absent"}
JOURS = ["lun.", "mar.", "mer.", "jeu.", "ven.", "sam.", "dim."]
POLICE = "Arial"
fin = Side(style="thin", color="D0D0D0")
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)
COULEURS = (("Présent", "E8F5E9", "2E7D32"), ("Excusé", "FFF1DC", "B35F00"),
            ("Absent", "FDE2E2", "C62828"))


def _nom_onglet(label, deja):
    """Nom d'onglet valide pour Excel (31 car. max, sans []:*?/\\), unique."""
    base = re.sub(r"[\[\]:*?/\\]", "-", label).strip()[:31] or "Cours"
    nom, i = base, 2
    while nom.lower() in deja:
        suffixe = f" ({i})"
        nom = base[: 31 - len(suffixe)] + suffixe
        i += 1
    deja.add(nom.lower())
    return nom


def _tri(ligne):
    return (ligne.get("essai", False), ligne["nom"].upper(), (ligne.get("prenom") or "").upper())


def _remplir_onglet(ws, cours):
    seances = sorted(cours["seances"], key=lambda s: (s["date"], s.get("horaire") or ""))
    lignes = sorted(cours["lignes"], key=_tri)
    n = len(seances)

    ws["A1"] = cours["label"]
    ws["A1"].font = Font(name=POLICE, bold=True, size=14)
    ws["A2"] = ("Case vide : pas sur la liste ce jour-là (non comptée). "
                "Totaux et taux calculés automatiquement.")
    ws["A2"].font = Font(name=POLICE, italic=True, size=9, color="777777")

    L_ENTETE, L1 = 4, 5
    c_tot = 3 + n
    titres = []
    for s in seances:
        d = s["date"]
        t = f"{JOURS[d.weekday()]} {d:%d/%m}"
        if s.get("horaire"):
            t += f"\n{s['horaire']}"
        titres.append(t)
    entetes = ["Nom", "Prénom"] + titres + ["Présences", "Absences", "Excusés", "Taux de présence"]
    for col, txt in enumerate(entetes, start=1):
        c = ws.cell(row=L_ENTETE, column=col, value=txt)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="37474F")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORD
    ws.row_dimensions[L_ENTETE].height = 34

    for i, ligne in enumerate(lignes):
        r = L1 + i
        ws.cell(row=r, column=1, value=ligne["nom"].upper())
        prenom = ligne.get("prenom") or ""
        if ligne.get("essai"):
            prenom = (prenom + " (essai)").strip()
        ws.cell(row=r, column=2, value=prenom)
        for j, s in enumerate(seances):
            st = ligne["statuts"].get(s["cle"])
            if st:
                ws.cell(row=r, column=3 + j, value=LIBELLES.get(st, "Absent"))
        if n:
            plage = f"C{r}:{get_column_letter(2 + n)}{r}"
            p = get_column_letter(c_tot)
            ws.cell(row=r, column=c_tot, value=f'=COUNTIF({plage},"Présent")')
            ws.cell(row=r, column=c_tot + 1, value=f'=COUNTIF({plage},"Absent")')
            ws.cell(row=r, column=c_tot + 2, value=f'=COUNTIF({plage},"Excusé")')
            ws.cell(row=r, column=c_tot + 3, value=f"=IFERROR({p}{r}/COUNTA({plage}),0)")
        for col in range(1, c_tot + 4):
            c = ws.cell(row=r, column=col)
            c.font = Font(name=POLICE, bold=(col == 1), italic=bool(ligne.get("essai")))
            c.border = BORD
            if col >= 3:
                c.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=c_tot + 3).number_format = "0%"

    if lignes and n:
        r_fin = L1 + len(lignes) - 1
        r_tot = r_fin + 1
        ws.cell(row=r_tot, column=1, value="Présents par séance")
        ws.merge_cells(start_row=r_tot, start_column=1, end_row=r_tot, end_column=2)
        for j in range(n):
            col = get_column_letter(3 + j)
            ws.cell(row=r_tot, column=3 + j, value=f'=COUNTIF({col}{L1}:{col}{r_fin},"Présent")')
        for col in range(1, 3 + n):
            c = ws.cell(row=r_tot, column=col)
            c.font = Font(name=POLICE, bold=True)
            c.fill = PatternFill("solid", fgColor="ECEFF1")
            c.border = BORD
            c.alignment = Alignment(horizontal="center" if col >= 3 else "left")
        zone = f"C{L1}:{get_column_letter(2 + n)}{r_fin}"
        for valeur, fond, txt in COULEURS:
            ws.conditional_formatting.add(zone, CellIsRule(
                operator="equal", formula=[f'"{valeur}"'],
                fill=PatternFill("solid", fgColor=fond), font=Font(color=txt, bold=True)))
    elif not lignes:
        ws.cell(row=L1, column=1, value="Aucun adhérent pour ce cours.").font = Font(name=POLICE, italic=True)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 18
    for j in range(n):
        ws.column_dimensions[get_column_letter(3 + j)].width = 10
    for k, w in enumerate((11, 11, 10, 12)):
        ws.column_dimensions[get_column_letter(c_tot + k)].width = w
    ws.freeze_panes = ws.cell(row=L1, column=3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{L_ENTETE}:{L_ENTETE}"
    coul = (cours.get("couleur") or "").lstrip("#")
    if re.fullmatch(r"[0-9A-Fa-f]{6}", coul):
        ws.sheet_properties.tabColor = coul


def construire_classeur(cours):
    wb = Workbook()
    wb.remove(wb.active)
    deja = set()
    for c in sorted(cours, key=lambda c: c["label"].lower()):
        _remplir_onglet(wb.create_sheet(_nom_onglet(c["label"], deja)), c)
    if not wb.sheetnames:
        wb.create_sheet("Aucun appel")["A1"] = "Aucun appel enregistré sur cette période."
    return wb
