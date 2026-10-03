"""UGD Présences : appel des cours, lié au planning partagé.

Les cours et créneaux viennent du planning (Google Sheet via Apps Script).
L'admin importe seulement les listes d'adhérents par cours.
Les entraîneurs choisissent leur nom, voient leurs cours du jour et font l'appel.
"""
import csv
import hashlib
import io
import json
import os
import re
import secrets
import time
import unicodedata
import urllib.request
from datetime import date, datetime, timedelta
from functools import wraps
from zoneinfo import ZoneInfo

from flask import (Flask, Response, abort, flash, redirect, render_template,
                   request, session, url_for)
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import inspect as sa_inspect, text as sa_text

# ------------------------------------------------------------------ config
TZ = ZoneInfo("Europe/Paris")
PLANNING_URL = os.environ.get(
    "PLANNING_URL",
    "https://script.google.com/macros/s/AKfycbxdgPFiy7ga4BTinEpuEbWbYBLBrfB8UgWeCTcN0ND-1rN9cB81vbStKwchPjM9DxRpLQ/exec",
)
CODE_ENTRAINEUR = os.environ.get("CODE_ENTRAINEUR", "").strip()
CODE_ADMIN = os.environ.get("CODE_ADMIN", "").strip()
CACHE_SECONDES = 300

JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
        "août", "septembre", "octobre", "novembre", "décembre"]

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY") or "dev-a-changer"
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=180)
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024

db_url = os.environ.get("DATABASE_URL", "sqlite:///presences.db")
for prefix in ("postgres://", "postgresql://"):
    if db_url.startswith(prefix):
        db_url = "postgresql+psycopg://" + db_url[len(prefix):]
app.config["SQLALCHEMY_DATABASE_URI"] = db_url
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True}
db = SQLAlchemy(app)


# ------------------------------------------------------------------ modèles
class Membre(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    groupe = db.Column(db.String(300), index=True, nullable=False)
    nom = db.Column(db.String(150), nullable=False)
    prenom = db.Column(db.String(150), default="")
    telephone = db.Column(db.String(40), default="")


class Appel(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    course_id = db.Column(db.String(80), nullable=False)
    jour_date = db.Column(db.Date, nullable=False, index=True)
    groupe = db.Column(db.String(300), index=True, nullable=False)
    intitule = db.Column(db.String(200))
    precision = db.Column(db.String(200))
    horaire = db.Column(db.String(20))
    lieu = db.Column(db.String(200))
    fait_par = db.Column(db.String(150))
    maj = db.Column(db.DateTime)
    presences = db.relationship("Presence", backref="appel",
                                cascade="all, delete-orphan")
    __table_args__ = (db.UniqueConstraint("course_id", "jour_date"),)


class Presence(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    appel_id = db.Column(db.Integer, db.ForeignKey("appel.id"), nullable=False, index=True)
    nom = db.Column(db.String(150), nullable=False)
    prenom = db.Column(db.String(150), default="")
    present = db.Column(db.Boolean, default=False)
    excuse = db.Column(db.Boolean, default=False)
    essai = db.Column(db.Boolean, default=False)


class Cache(db.Model):
    cle = db.Column(db.String(50), primary_key=True)
    valeur = db.Column(db.Text)
    maj = db.Column(db.DateTime)


with app.app_context():
    db.create_all()
    # Migration douce : ajoute la colonne téléphone aux bases déjà existantes
    if "telephone" not in [c["name"] for c in sa_inspect(db.engine).get_columns("membre")]:
        with db.engine.begin() as cx:
            cx.execute(sa_text("ALTER TABLE membre ADD COLUMN telephone VARCHAR(40) DEFAULT ''"))
    # … et la colonne « excusé » aux présences
    if "excuse" not in [c["name"] for c in sa_inspect(db.engine).get_columns("presence")]:
        with db.engine.begin() as cx:
            cx.execute(sa_text("ALTER TABLE presence ADD COLUMN excuse BOOLEAN DEFAULT FALSE"))


# ------------------------------------------------------------------ outils
def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().replace("’", "'").split())


def cle_groupe(intitule, precision):
    return norm(intitule) + "|" + norm(precision)


def gid(cle):
    return hashlib.sha1(cle.encode()).hexdigest()[:10]


def cle_personne(nom, prenom):
    return norm(nom) + "|" + norm(prenom)


def format_tel(v):
    """Normalise un numéro : '612345678', '+33 6 12…', '0612345678.0' -> '06 12 34 56 78'."""
    brut = str(v or "").strip()
    if brut.endswith(".0"):
        brut = brut[:-2]
    chiffres = "".join(c for c in brut if c.isdigit())
    if chiffres.startswith("0033"):
        chiffres = "0" + chiffres[4:]
    elif chiffres.startswith("33") and len(chiffres) == 11:
        chiffres = "0" + chiffres[2:]
    elif len(chiffres) == 9 and not chiffres.startswith("0"):
        chiffres = "0" + chiffres  # zéro perdu par Excel
    if len(chiffres) == 10 and chiffres.startswith("0"):
        return " ".join(chiffres[i:i + 2] for i in range(0, 10, 2))
    return brut[:40]


def maintenant():
    return datetime.now(TZ)


def date_fr(d):
    return f"{JOURS[d.weekday()]} {d.day} {MOIS[d.month - 1]}"


app.jinja_env.filters["date_fr"] = date_fr
app.jinja_env.globals["gid"] = gid


# ------------------------------------------------------------------ planning
_planning = {"t": 0.0, "data": None, "erreur": None}


def _lire_distant():
    url = PLANNING_URL + ("&" if "?" in PLANNING_URL else "?") + "t=" + str(int(time.time()))
    req = urllib.request.Request(url, headers={"User-Agent": "UGD-Presences"})
    with urllib.request.urlopen(req, timeout=20) as r:
        return json.loads(r.read().decode("utf-8"))


def _analyser(brut):
    reglages = brut.get("settings") or {}
    couleurs = {}
    for d in reglages.get("discs") or []:
        couleurs[norm(d.get("name"))] = d.get("color")
        couleurs[norm(d.get("id"))] = d.get("color")
    cours = []
    for i, r in enumerate(brut.get("courses") or []):
        intitule = str(r.get("intitule") or "").strip()
        if not intitule:
            continue
        jour_txt = norm(r.get("jour"))
        jour = next((k for k, j in enumerate(JOURS) if norm(j) == jour_txt), None)
        if jour is None:
            continue
        precision = str(r.get("precision") or "").strip()
        cle = cle_groupe(intitule, precision)
        cours.append({
            "id": str(r.get("id") or f"x{i}"),
            "intitule": intitule,
            "precision": precision,
            "label": intitule + (f" – {precision}" if precision else ""),
            "groupe": cle,
            "gid": gid(cle),
            "discipline": str(r.get("discipline") or ""),
            "couleur": couleurs.get(norm(r.get("discipline"))) or "#5B6475",
            "jour": jour,
            "debut": str(r.get("debut") or "00:00")[:5],
            "fin": str(r.get("fin") or "00:00")[:5],
            "lieu": str(r.get("lieu") or ""),
            "espace": str(r.get("espace") or ""),
            "encadrants": [x.strip() for x in str(r.get("encadrants") or "").split("/") if x.strip()],
        })
    cours.sort(key=lambda c: (c["jour"], c["debut"], c["label"]))
    return {"cours": cours}


def planning(forcer=False):
    """Planning en cache mémoire 5 min, avec copie de secours en base."""
    if not forcer and _planning["data"] and time.time() - _planning["t"] < CACHE_SECONDES:
        return _planning["data"]
    try:
        data = _analyser(_lire_distant())
        if not data["cours"]:
            raise ValueError("le planning partagé est vide")
        _planning.update(t=time.time(), data=data, erreur=None)
        ligne = db.session.get(Cache, "planning") or Cache(cle="planning")
        ligne.valeur, ligne.maj = json.dumps(data), datetime.utcnow()
        db.session.add(ligne)
        db.session.commit()
    except Exception as e:  # réseau, script Google indisponible…
        app.logger.warning("Lecture du planning impossible : %s", e)
        _planning["erreur"] = str(e)
        if not _planning["data"]:
            ligne = db.session.get(Cache, "planning")
            _planning["data"] = json.loads(ligne.valeur) if ligne else {"cours": []}
        _planning["t"] = time.time() - CACHE_SECONDES + 60  # on réessaie dans 1 min
    return _planning["data"]


def groupes():
    res = {}
    for c in planning()["cours"]:
        g = res.setdefault(c["groupe"], {"cle": c["groupe"], "gid": c["gid"], "label": c["label"],
                                         "couleur": c["couleur"], "creneaux": []})
        g["creneaux"].append(c)
    return sorted(res.values(), key=lambda g: norm(g["label"]))


def groupe_par_gid(g):
    return next((x for x in groupes() if x["gid"] == g), None)


def encadrants():
    noms = {}
    for c in planning()["cours"]:
        for n in c["encadrants"]:
            noms.setdefault(norm(n), n)
    return sorted(noms.values(), key=norm)


# ------------------------------------------------------------------ sécurité
@app.before_request
def csrf():
    session.permanent = True
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    if request.method == "POST" and request.form.get("csrf") != session["csrf"]:
        abort(400)


app.jinja_env.globals["csrf"] = lambda: session.get("csrf", "")


def connecte(f):
    @wraps(f)
    def w(*a, **k):
        if session.get("role") not in ("coach", "admin"):
            return redirect(url_for("connexion", suite=request.full_path))
        return f(*a, **k)
    return w


def admin(f):
    @wraps(f)
    def w(*a, **k):
        if session.get("role") != "admin":
            return redirect(url_for("connexion", suite=request.full_path, admin=1))
        return f(*a, **k)
    return w


# ------------------------------------------------------------------ connexion
@app.route("/connexion", methods=["GET", "POST"])
def connexion():
    erreur = None
    if request.method == "POST":
        code = request.form.get("code", "").strip()
        if CODE_ADMIN and code == CODE_ADMIN:
            session["role"] = "admin"
        elif CODE_ENTRAINEUR and code == CODE_ENTRAINEUR:
            session["role"] = "coach"
        else:
            erreur = "Code incorrect."
        if not erreur:
            suite = request.args.get("suite") or url_for("accueil")
            return redirect(suite if suite.startswith("/") else url_for("accueil"))
    non_configure = not CODE_ENTRAINEUR and not CODE_ADMIN
    return render_template("connexion.html", erreur=erreur, non_configure=non_configure)


@app.route("/deconnexion")
def deconnexion():
    session.clear()
    return redirect(url_for("connexion"))


@app.route("/qui", methods=["GET", "POST"])
@connecte
def qui():
    if request.method == "POST":
        session["coach"] = request.form.get("nom", "")
        return redirect(url_for("accueil"))
    return render_template("qui.html", noms=encadrants(), actuel=session.get("coach"))


# ------------------------------------------------------------------ entraîneurs
@app.route("/")
@connecte
def accueil():
    if "coach" not in session:
        return redirect(url_for("qui"))
    now = maintenant()
    try:
        jour = date.fromisoformat(request.args.get("d", ""))
    except ValueError:
        jour = now.date()
    coach = session.get("coach") or ""
    tous = request.args.get("tous") == "1" or not coach

    du_jour = [c for c in planning()["cours"] if c["jour"] == jour.weekday()]
    miens = [c for c in du_jour if norm(coach) in {norm(n) for n in c["encadrants"]}]
    liste = du_jour if tous else miens

    faits = {a.course_id: a for a in Appel.query.filter_by(jour_date=jour).all()}
    heure = now.strftime("%H:%M")
    cartes = []
    for c in liste:
        etat = "avenir"
        if jour < now.date() or (jour == now.date() and c["fin"] <= heure):
            etat = "passe"
        elif jour == now.date() and c["debut"] <= heure < c["fin"]:
            etat = "encours"
        a = faits.get(c["id"])
        nb = (sum(p.present for p in a.presences), len(a.presences)) if a else None
        cartes.append({**c, "etat": etat, "fait": nb})
    ordre = {"encours": 0, "avenir": 1, "passe": 2}
    if jour == now.date():
        cartes.sort(key=lambda c: (ordre[c["etat"]], c["debut"]))

    return render_template("accueil.html", cartes=cartes, jour=jour, coach=coach, tous=tous,
                           nb_miens=len(miens), nb_jour=len(du_jour),
                           hier=jour - timedelta(days=1), demain=jour + timedelta(days=1),
                           aujourdhui=now.date(), erreur_planning=_planning["erreur"])


@app.route("/appel/<course_id>", methods=["GET", "POST"])
@connecte
def appel(course_id):
    c = next((x for x in planning()["cours"] if x["id"] == course_id), None)
    if not c:
        abort(404)
    try:
        jour = date.fromisoformat(request.args.get("d", ""))
    except ValueError:
        jour = maintenant().date()
    membres = Membre.query.filter_by(groupe=c["groupe"]).all()
    membres.sort(key=lambda m: (norm(m.nom), norm(m.prenom)))
    existant = Appel.query.filter_by(course_id=c["id"], jour_date=jour).first()

    if request.method == "POST":
        presents = set(request.form.getlist("present"))
        excuses = set(request.form.getlist("excuse")) - presents  # « présent » l'emporte
        if not existant:
            existant = Appel(course_id=c["id"], jour_date=jour)
            db.session.add(existant)
        existant.groupe, existant.intitule, existant.precision = c["groupe"], c["intitule"], c["precision"]
        existant.horaire, existant.lieu = f'{c["debut"]}-{c["fin"]}', c["lieu"]
        existant.fait_par = session.get("coach") or ("admin" if session.get("role") == "admin" else "")
        existant.maj = datetime.utcnow()
        existant.presences.clear()
        for m in membres:
            existant.presences.append(Presence(nom=m.nom, prenom=m.prenom,
                                               present=str(m.id) in presents,
                                               excuse=str(m.id) in excuses))
        anciens_essais = request.form.getlist("essai_garde")
        nouveaux = [x.strip() for x in request.form.get("essais", "").replace(";", ",").split(",")]
        for n in anciens_essais + nouveaux:
            if n.strip():
                existant.presences.append(Presence(nom=n.strip(), present=True, essai=True))
        db.session.commit()
        nb = sum(p.present for p in existant.presences)
        nb_exc = sum(bool(p.excuse) for p in existant.presences)
        flash(f"Appel enregistré : {nb} présent{'s' if nb > 1 else ''} sur {len(existant.presences)}"
              + (f", {nb_exc} excusé{'s' if nb_exc > 1 else ''}" if nb_exc else "") + ".")
        return redirect(url_for("accueil", d=jour.isoformat()) if jour != maintenant().date()
                        else url_for("accueil"))

    coches, excuses, essais = set(), set(), []
    if existant:
        deja = {cle_personne(p.nom, p.prenom): p for p in existant.presences if not p.essai}
        for m in membres:
            p = deja.get(cle_personne(m.nom, m.prenom))
            if p and p.present:
                coches.add(m.id)
            elif p and p.excuse:
                excuses.add(m.id)
        essais = [p.nom for p in existant.presences if p.essai]
    return render_template("appel.html", c=c, jour=jour, membres=membres, coches=coches,
                           excuses=excuses, essais=essais, existant=existant)


# ------------------------------------------------------------------ admin
@app.route("/admin", methods=["GET", "POST"])
@admin
def admin_accueil():
    if request.method == "POST" and request.form.get("action") == "rafraichir":
        planning(forcer=True)
        if _planning["erreur"]:
            flash("Le planning n'a pas pu être relu : " + _planning["erreur"], "err")
        else:
            flash("Planning relu.")
        return redirect(url_for("admin_accueil"))
    gs = groupes()
    nb_membres = dict(db.session.query(Membre.groupe, db.func.count(Membre.id)).group_by(Membre.groupe).all())
    nb_appels = dict(db.session.query(Appel.groupe, db.func.count(Appel.id)).group_by(Appel.groupe).all())
    for g in gs:
        g["nb_membres"] = nb_membres.get(g["cle"], 0)
        g["nb_appels"] = nb_appels.get(g["cle"], 0)
    return render_template("admin.html", groupes=gs, erreur_planning=_planning["erreur"])


def _lire_feuilles(f):
    """Renvoie [(nom_onglet, lignes)] ; un fichier CSV/texte donne un seul onglet sans nom."""
    nom = (f.filename or "").lower()
    contenu = f.read()
    if contenu[:4] == b"\xd0\xcf\x11\xe0" or nom.endswith(".xls"):
        # Excel 97-2003 (.xls), format des exports Kalisport
        import xlrd
        wb = xlrd.open_workbook(file_contents=contenu)
        return [(sh.name, [["" if v is None else str(v).strip() for v in sh.row_values(r)]
                           for r in range(sh.nrows)])
                for sh in wb.sheets()]
    if contenu[:2] == b"PK" or nom.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        return [(ws.title, [["" if v is None else str(v).strip() for v in row]
                            for row in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets]
    texte = ""
    for enc in ("utf-8-sig", "cp1252"):
        try:
            texte = contenu.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return [("", _lire_texte(texte))]


def compact(s):
    """Clé de comparaison tolérante : sans accents, espaces, tirets ni ponctuation."""
    return re.sub(r"[^a-z0-9]", "", norm(s))


def _trouver_groupe(texte, gs):
    """Associe un nom de cours (colonne ou onglet) à un cours du planning."""
    k = compact(texte)
    if not k:
        return None
    for g in gs:
        if compact(g["label"]) == k:
            return g
    meme = [g for g in gs if compact(g["creneaux"][0]["intitule"]) == k]
    if len(meme) == 1:
        return meme[0]
    if len(str(texte)) >= 28:  # nom d'onglet tronqué par Excel (31 caractères max)
        debut = [g for g in gs if compact(g["label"]).startswith(k)]
        if len(debut) == 1:
            return debut[0]
    return None


def _lire_texte(texte):
    lignes = [l for l in texte.splitlines() if l.strip()]
    if not lignes:
        return []
    sep = max(["\t", ";", ","], key=lambda s: lignes[0].count(s))
    if lignes[0].count(sep) == 0:
        return [[l.strip()] for l in lignes]
    return [[x.strip() for x in r] for r in csv.reader(lignes, delimiter=sep)]


def _colonnes(lignes):
    """Repère les colonnes nom / prénom / cours / téléphone ;
    sinon nom=1re, prénom=2e, téléphone=3e."""
    if not lignes:
        return [], {}
    entete = [norm(x) for x in lignes[0]]
    idx = {}
    tels = [i for i, h in enumerate(entete)
            if any(k in h for k in ("tel", "portable", "mobile", "gsm"))]
    if tels:
        # on préfère un portable s'il y a plusieurs colonnes téléphone
        idx["tel"] = next((i for i in tels if any(k in entete[i] for k in ("portable", "mobile", "gsm"))),
                          tels[0])
    for i, h in enumerate(entete):
        if i in tels:
            continue
        if "prenom" in h and "prenom" not in idx:
            idx["prenom"] = i
        elif (h in ("nom", "noms") or re.match(r"nom\W", h)) and "nom" not in idx:
            idx["nom"] = i
        elif any(k in h for k in ("cours", "groupe", "activite", "section")) and "cours" not in idx:
            idx["cours"] = i
    if "nom" in idx:
        return lignes[1:], idx
    idx = {"nom": 0, "prenom": 1}
    if any(len(r) > 2 for r in lignes):
        idx["tel"] = 2
    return lignes, idx


def _personnes(lignes, idx):
    res = []
    for r in lignes:
        nom = r[idx["nom"]].strip() if len(r) > idx["nom"] else ""
        prenom = r[idx["prenom"]].strip() if "prenom" in idx and len(r) > idx["prenom"] else ""
        if nom:
            cours = r[idx["cours"]].strip() if "cours" in idx and len(r) > idx["cours"] else ""
            tel = format_tel(r[idx["tel"]]) if "tel" in idx and len(r) > idx["tel"] else ""
            res.append((nom, prenom, cours, tel))
    return res


def _enregistrer_liste(cle, personnes, remplacer):
    existants = Membre.query.filter_by(groupe=cle).all()
    par_cle = {cle_personne(m.nom, m.prenom): m for m in existants}
    voulus = {}
    for nom, prenom, _, tel in personnes:
        voulus.setdefault(cle_personne(nom, prenom), (nom, prenom, tel))
    ajout = 0
    for k, (nom, prenom, tel) in voulus.items():
        if k not in par_cle:
            db.session.add(Membre(groupe=cle, nom=nom, prenom=prenom, telephone=tel))
            ajout += 1
        elif tel:
            par_cle[k].telephone = tel  # met à jour le numéro d'un adhérent déjà inscrit
    retrait = 0
    if remplacer:
        for k, m in par_cle.items():
            if k not in voulus:
                db.session.delete(m)
                retrait += 1
    db.session.commit()
    return ajout, retrait


@app.route("/admin/groupe/<g>", methods=["GET", "POST"])
@admin
def admin_groupe(g):
    grp = groupe_par_gid(g)
    if not grp:
        abort(404)
    if request.method == "POST":
        action = request.form.get("action")
        if action == "importer":
            f = request.files.get("fichier")
            personnes, onglets = None, []
            if f and f.filename:
                feuilles = _lire_feuilles(f)
                onglets = [n for n, _ in feuilles]
                if len(feuilles) == 1:
                    lignes = feuilles[0][1]
                else:  # classeur à plusieurs onglets : on prend celui de ce cours
                    lignes = next((l for n, l in feuilles if _trouver_groupe(n, [grp])), None)
            else:
                lignes = _lire_texte(request.form.get("texte", ""))
            if lignes is not None:
                personnes = _personnes(*_colonnes(lignes))
            if lignes is None:
                flash(f"Aucun onglet du fichier ne correspond à « {grp['label']} ». "
                      f"Onglets trouvés : {', '.join(onglets)}. Renommez l'onglet comme le cours.", "err")
            elif not personnes:
                flash("Aucun nom trouvé. Collez une personne par ligne (nom, prénom, téléphone) "
                      "ou choisissez un fichier.", "err")
            else:
                a, r = _enregistrer_liste(grp["cle"], personnes, request.form.get("mode") == "remplacer")
                flash(f"Liste mise à jour : {a} ajout(s), {r} retrait(s).")
        elif action == "ajouter":
            nom = request.form.get("nom", "").strip()
            if nom:
                _enregistrer_liste(grp["cle"], [(nom, request.form.get("prenom", "").strip(), "",
                                                 format_tel(request.form.get("telephone", "")))], False)
                flash("Adhérent ajouté.")
        elif action == "telephone":
            m = db.session.get(Membre, int(request.form.get("id", 0)))
            if m and m.groupe == grp["cle"]:
                m.telephone = format_tel(request.form.get("telephone", ""))
                db.session.commit()
                flash(f"Téléphone de {m.nom} {m.prenom} mis à jour.")
        elif action == "retirer":
            m = db.session.get(Membre, int(request.form.get("id", 0)))
            if m and m.groupe == grp["cle"]:
                db.session.delete(m)
                db.session.commit()
                flash(f"{m.nom} {m.prenom} retiré de la liste.")
        return redirect(url_for("admin_groupe", g=g))

    membres = sorted(Membre.query.filter_by(groupe=grp["cle"]).all(), key=lambda m: (norm(m.nom), norm(m.prenom)))
    appels = Appel.query.filter_by(groupe=grp["cle"]).order_by(Appel.jour_date.desc(), Appel.horaire.desc()).limit(16).all()
    appels.reverse()
    grille = {}
    for a in appels:
        for p in a.presences:
            grille[(cle_personne(p.nom, p.prenom), a.id)] = (
                "present" if p.present else "excuse" if p.excuse else "absent")
    tous_appels = Appel.query.filter_by(groupe=grp["cle"]).all()
    totaux = {}
    for a in tous_appels:
        for p in a.presences:
            t = totaux.setdefault(cle_personne(p.nom, p.prenom), [0, 0, 0])  # présent, séances, excusé
            t[0] += bool(p.present)
            t[1] += 1
            t[2] += bool(p.excuse) and not p.present
    lignes = [{"m": m, "k": cle_personne(m.nom, m.prenom),
               "total": totaux.get(cle_personne(m.nom, m.prenom), [0, 0, 0])} for m in membres]
    return render_template("groupe.html", grp=grp, lignes=lignes, appels=appels, grille=grille,
                           nb_appels=len(tous_appels))


@app.route("/admin/import", methods=["POST"])
@admin
def admin_import():
    """Import global : soit une colonne cours/groupe, soit un onglet par cours
    (export Kalisport : onglet = cours, lignes = nom, prénom, téléphone)."""
    f = request.files.get("fichier")
    if not f or not f.filename:
        flash("Choisissez un fichier.", "err")
        return redirect(url_for("admin_accueil"))
    try:
        feuilles = _lire_feuilles(f)
    except Exception as e:
        app.logger.warning("Fichier illisible : %s", e)
        flash("Fichier illisible. Enregistrez-le en .xlsx, .xls ou .csv et réessayez.", "err")
        return redirect(url_for("admin_accueil"))
    gs = groupes()
    lots, inconnus = {}, {}
    for onglet, brut in feuilles:
        lignes, idx = _colonnes(brut)
        personnes = _personnes(lignes, idx)
        if not personnes:
            continue
        if "cours" in idx:  # une colonne indique le cours de chaque ligne
            for p in personnes:
                g = _trouver_groupe(p[2], gs)
                if g:
                    lots.setdefault(g["cle"], []).append(p)
                else:
                    inconnus[p[2] or "(vide)"] = inconnus.get(p[2] or "(vide)", 0) + 1
        else:  # sinon, c'est le nom de l'onglet qui donne le cours
            g = _trouver_groupe(onglet, gs)
            if g:
                lots.setdefault(g["cle"], []).extend(personnes)
            else:
                inconnus[onglet or "(fichier sans colonne cours)"] = len(personnes)
    total = 0
    for cle, personnes in lots.items():
        total += _enregistrer_liste(cle, personnes, request.form.get("mode") == "remplacer")[0]
    flash(f"{total} adhérent(s) ajouté(s) dans {len(lots)} cours.")
    if inconnus:
        detail = ", ".join(f"« {k} » ({v})" for k, v in sorted(inconnus.items()))
        flash("Cours non reconnus dans le planning, lignes ignorées : " + detail +
              ". Renommez l'onglet (ou la valeur de la colonne cours) exactement comme le cours "
              "dans le planning, puis réimportez.", "err")
    return redirect(url_for("admin_accueil"))


@app.route("/admin/export.csv")
@admin
def admin_export():
    q = Appel.query
    g = request.args.get("g")
    if g:
        grp = groupe_par_gid(g)
        if grp:
            q = q.filter_by(groupe=grp["cle"])
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["Date", "Jour", "Horaire", "Cours", "Précision", "Lieu", "Appel fait par",
                "Nom", "Prénom", "Présent", "Excusé", "Essai"])
    for a in q.order_by(Appel.jour_date, Appel.horaire).all():
        for p in sorted(a.presences, key=lambda p: (norm(p.nom), norm(p.prenom))):
            w.writerow([a.jour_date.strftime("%d/%m/%Y"), JOURS[a.jour_date.weekday()], a.horaire,
                        a.intitule, a.precision, a.lieu, a.fait_par, p.nom, p.prenom,
                        "oui" if p.present else "non", "oui" if p.excuse else "",
                        "oui" if p.essai else ""])
    return Response("\ufeff" + out.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=presences.csv"})


def _periode():
    """Filtre optionnel ?du=AAAA-MM-JJ&au=AAAA-MM-JJ."""
    res = []
    for k in ("du", "au"):
        try:
            res.append(date.fromisoformat(request.args.get(k, "")))
        except ValueError:
            res.append(None)
    return res


@app.route("/admin/export.xlsx")
@admin
def admin_export_xlsx():
    """Classeur Excel : un onglet par cours, une ligne par adhérent, une colonne par séance."""
    from export_presences import construire_classeur

    q = Appel.query
    g = request.args.get("g")
    if g:
        grp = groupe_par_gid(g)
        if grp:
            q = q.filter_by(groupe=grp["cle"])
    du, au = _periode()
    if du:
        q = q.filter(Appel.jour_date >= du)
    if au:
        q = q.filter(Appel.jour_date <= au)
    appels = q.order_by(Appel.jour_date, Appel.horaire).all()

    infos = {x["cle"]: x for x in groupes()}
    par_groupe = {}
    for a in appels:
        par_groupe.setdefault(a.groupe, []).append(a)

    cours = []
    for cle, liste in par_groupe.items():
        info = infos.get(cle)
        label = info["label"] if info else (liste[0].intitule or "Cours") + (
            f" – {liste[0].precision}" if liste[0].precision else "")
        # Deux séances le même jour (deux créneaux) : on précise l'horaire dans l'en-tête
        par_jour = {}
        for a in liste:
            par_jour[a.jour_date] = par_jour.get(a.jour_date, 0) + 1
        seances = [{"cle": a.id, "date": a.jour_date,
                    "horaire": a.horaire if par_jour[a.jour_date] > 1 else ""} for a in liste]

        lignes = {}
        # adhérents actuellement inscrits (même s'ils n'apparaissent dans aucun appel)
        for m in Membre.query.filter_by(groupe=cle).all():
            lignes[cle_personne(m.nom, m.prenom)] = {"nom": m.nom, "prenom": m.prenom,
                                                     "essai": False, "statuts": {}}
        for a in liste:
            for p in a.presences:
                k = ("essai|" if p.essai else "") + cle_personne(p.nom, p.prenom)
                ligne = lignes.setdefault(k, {"nom": p.nom, "prenom": p.prenom,
                                              "essai": p.essai, "statuts": {}})
                ligne["statuts"][a.id] = ("present" if p.present
                                          else "excuse" if p.excuse else "absent")
        cours.append({"label": label, "couleur": info["couleur"] if info else None,
                      "seances": seances, "lignes": list(lignes.values())})

    buf = io.BytesIO()
    construire_classeur(cours).save(buf)
    buf.seek(0)
    nom = "presences-ugd" + (f"-{du:%Y%m%d}" if du else "") + (f"-{au:%Y%m%d}" if au else "") + ".xlsx"
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={nom}"})


# ------------------------------------------------------------------ erreurs
@app.errorhandler(404)
def e404(e):
    return render_template("erreur.html", titre="Page introuvable",
                           texte="Cette page n'existe pas ou plus."), 404


@app.errorhandler(400)
def e400(e):
    return render_template("erreur.html", titre="Formulaire expiré",
                           texte="Rechargez la page et recommencez."), 400


@app.errorhandler(500)
def e500(e):
    return render_template("erreur.html", titre="Erreur",
                           texte="Un problème est survenu. Réessayez dans un instant."), 500


if __name__ == "__main__":
    app.run(debug=True)
